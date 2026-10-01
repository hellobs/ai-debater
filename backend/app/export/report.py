"""复盘导出：Markdown / HTML（打印成 PDF）/ Word(.docx)。

为什么 PDF 走"打印优化 HTML"而不是直接生成 PDF
------------------------------------------------
在 macOS/浏览器环境下，直接生成中文 PDF 需要内嵌 CJK 字体（fpdf2/reportlab 都得
挂一个 TTF），字体缺失就会变成一排方块。而浏览器"打印 → 另存为 PDF"用的是系统字体，
中文排版最好、零依赖、零字体风险。

所以：Markdown 与 Word 由后端直接产出文件；PDF 由后端产出一张**打印优化页面**，
前端一键唤起打印对话框（用户在对话框里选"存储为 PDF"）。
"""
from __future__ import annotations

import html
import io
import logging
from datetime import datetime
from typing import Any

from ..ledger import store

logger = logging.getLogger("export")

#: 名册读不出来时的兜底名表。**正常路径不读它** ——
#: 显示名的唯一来源是 `configs/advisors.yaml`（见 `advisor_titles()`）。
#: 兜底存在的理由：导出**不该因为名册问题变成一份没有正文的空报告**，
#: 那比印错一个字严重得多。
_FALLBACK_TITLES = {
    "rebutter": "反驳手",
    "questioner": "质询手",
    "auditor": "逻辑审计员",
    "strategist": "论证策略师",
    "risk": "风险提示员",
}

#: 底座出处。本项目的基础设施半边（模型接入 / 提示词模板 / 插件总线）**基于 mavis 开发**，
#: 且是**只读依赖**——导出物里如实署名，避免读者以为这些是自研的。
MAVIS_NAME = "mavis"
MAVIS_HOME = "https://github.com/hellobs/mavis"

#: 归属行末尾那句固定说明。三个出口（Markdown / HTML / Word）共用，改一处即三处生效。
BASE_NOTE = "底座能力：模型接入 / 提示词模板 / 插件总线"

#: 导出报告「使用提示」的固定条目。**必须领域中立** ——
#: 导出是**平台级**产物，与辩题落在哪个领域无关。这一条此前写的是
#: 「标注「待核验」的**法源**引用尚未经过引用回链核验，**上庭前**请自行确认」，
#: 于是每一份**通用**辩题（例如"大学应否把 AI 设为必修课"）的复盘里都带着法庭措辞。
#: 这是提示词 / schema 描述 / 显示名之外的**第四处**领域泄漏，且直接出现在交付物上。
NOTE_UNVERIFIED = "标注「待核验」的引用尚未经过回链核验，正式使用前请自行确认。"
#: 同一条的带格式变体。**从上面派生**，免得两个出口的措辞各改一半。
NOTE_UNVERIFIED_MD = NOTE_UNVERIFIED.replace("「待核验」", "「**待核验**」")
NOTE_UNVERIFIED_HTML = NOTE_UNVERIFIED.replace("「待核验」", "「<b>待核验</b>」")


def advisor_titles() -> dict[str, str]:
    """`{参谋名: 显示名}`，**顺序即导出章节顺序**。

    显示名取自名册 —— `configs/advisors.yaml` 是唯一来源，这里不另写一份。
    此前这里是写死的字典，实测代价：`advisors.yaml` 把 strategist 的中文名改成
    领域中立的「论证策略师」之后，**导出报告仍然印着「解释方法策略师」**；
    同一个显示名被定义两遍，改一处不算改完。顺序也直接沿用名册顺序（＝界面列序）。
    """
    titles: dict[str, str] = {}
    try:
        from ..advisors import load_roster

        titles = {a.name: (a.label or a.name) for a in load_roster()}
    except Exception:  # noqa: BLE001
        logger.warning("取名册显示名失败，导出退回内置兜底名表", exc_info=True)
    if not titles:
        titles = dict(_FALLBACK_TITLES)
    for name, label in _FALLBACK_TITLES.items():
        titles.setdefault(name, label)      # 旧会话里可能有名册已删掉的一路
    return titles


def _attribution(label: str) -> str:
    """归属声明。句式来自 `mavis_bridge.declaration()`，这里只把框架名换成带链接/粗体的形式。

    单一来源的意义：README、`/api/health`、导出报告说的是同一句话，
    不会出现"文档写基于 X 开发、交付物写成只用过 X"这种分裂。
    """
    try:
        from .. import mavis_bridge

        return mavis_bridge.declaration(label)
    except Exception:  # noqa: BLE001
        return f"本项目基于 {label} 开发（只读依赖，一行未改）"

KIND_TITLES = {
    "rebuttal": "反驳要点（四段：主张 / 大前提 / 小前提 / 结论）",
    "questions": "质询问题",
    "audit": "逻辑谬误指认",
    "strategy": "衡量尺度争夺点",
    "risk": "风险提示",
}


# --------------------------------------------------------------------------
# 数据汇总
# --------------------------------------------------------------------------
def build_report(session_id: str) -> dict | None:
    """把一个会话整理成导出用的结构化数据。"""
    session = store.get_session(session_id)
    if not session:
        return None

    turns = store.list_turns(session_id)
    cards = store.list_cards(session_id)
    suggestions = store.list_suggestions(session_id)

    # suggestions 是倒序存的，按 advisor 归并（同一路只保留最新一条）
    latest: dict[str, dict] = {}
    for s in suggestions:
        latest.setdefault(s["advisor"], s)

    sections = []
    for name, title in advisor_titles().items():
        s = latest.get(name)
        if not s:
            continue
        sections.append({
            "advisor": name,
            "title": title,
            "kind": s.get("status"),
            "latency_s": s.get("latency_s"),
            "payload": s.get("payload"),
        })

    return {
        "session": session,
        "turns": turns,
        "cards": cards,
        "sections": sections,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "stats": {
            "turns": len(turns),
            "cards": len(cards),
            "standing": sum(1 for c in cards if c["status"] == "standing"),
        },
    }


# --------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------
def _fmt_rebuttals(items: list[dict]) -> str:
    out = []
    for i, it in enumerate(items, 1):
        out.append(f"**{i}. {it.get('claim', '')}**\n")
        out.append(f"- 大前提：{it.get('major_premise', '')}")
        out.append(f"- 小前提：{it.get('minor_premise', '')}")
        out.append(f"- 结　论：{it.get('conclusion', '')}")
        out.append("")
    return "\n".join(out)


def to_markdown(data: dict) -> str:
    s = data["session"]
    lines: list[str] = []
    lines.append(f"# 辩论参谋复盘 — {s['topic']}")
    lines.append("")
    lines.append(f"- **我方立场**：{s['our_side']}")
    lines.append(f"- **会话编号**：`{s['id']}`")
    lines.append(f"- **生成时间**：{data['generated_at']}")
    lines.append(
        f"- **对方发言轮次**：{data['stats']['turns']}　"
        f"**台账条目**：{data['stats']['cards']}（其中成立 {data['stats']['standing']}）"
    )
    lines.append("")

    lines.append("## 一、对方发言记录")
    lines.append("")
    if not data["turns"]:
        lines.append("_（无）_")
    for i, t in enumerate(data["turns"], 1):
        lines.append(f"{i}. {t['opponent_text']}")
    lines.append("")

    lines.append("## 二、参谋建议")
    lines.append("")
    if not data["sections"]:
        lines.append("_（无）_")
    for sec in data["sections"]:
        lines.append(f"### {sec['title']}")
        lines.append("")
        payload = sec.get("payload")
        if sec["advisor"] == "rebutter" and isinstance(payload, list):
            lines.append(_fmt_rebuttals(payload))
        elif sec["advisor"] == "questioner" and isinstance(payload, list):
            for i, q in enumerate(payload, 1):
                lines.append(f"{i}. {q}")
            lines.append("")
        elif sec["advisor"] == "auditor" and isinstance(payload, list):
            if not payload:
                lines.append("_未发现明显逻辑谬误。_")
            for it in payload:
                if isinstance(it, dict):
                    lines.append(
                        f"- **{it.get('fallacy', '谬误')}**｜原话：「{it.get('quote', '')}」"
                    )
                    lines.append(f"  - {it.get('explain', '')}")
            lines.append("")
        elif sec["advisor"] == "strategist" and isinstance(payload, list):
            if not payload:
                lines.append("_（未识别到衡量尺度之争）_")
            for it in payload:
                if isinstance(it, dict):
                    lines.append(
                        f"- 对方用 **{it.get('opponent_method', '?')}**"
                        f"（{it.get('opponent_effect', '')}）"
                    )
                    lines.append(
                        f"  - 我方主张 **{it.get('our_method', '?')}** 优先："
                        f"{it.get('counter', '')}"
                    )
            lines.append("")
        elif sec["advisor"] == "risk" and isinstance(payload, list):
            if not payload:
                lines.append("_（未识别到明显风险）_")
            for it in payload:
                if isinstance(it, dict):
                    lines.append(f"- **[{it.get('kind', '风险')}]** {it.get('risk', '')}")
                    lines.append(f"  - 应对：{it.get('suggestion', '')}")
            lines.append("")

    lines.append("## 三、我方论点台账")
    lines.append("")
    if not data["cards"]:
        lines.append("_（无）_")
    else:
        lines.append("| 编号 | 主张 | 大前提/依据 | 状态 |")
        lines.append("|---|---|---|---|")
        for c in data["cards"]:
            status = {"standing": "成立", "weakened": "受损", "abandoned": "放弃"}.get(
                c["status"], c["status"]
            )
            lines.append(
                f"| {c['id']} | {c['claim']} | {c.get('major_premise') or '—'} | {status} |"
            )
    lines.append("")

    lines.append("## 四、使用提示")
    lines.append("")
    lines.append("- 以上内容由 AI 参谋生成，**仅作参考**，最终判断与取舍由你决定。")
    lines.append(f"- {NOTE_UNVERIFIED_MD}")
    lines.append("")
    lines.append("---")
    lines.append("")
    lines.append(
        f"{_attribution(f'**[{MAVIS_NAME}]({MAVIS_HOME})**')}"
        f"　｜　{BASE_NOTE}"
    )
    lines.append("")
    return "\n".join(lines)


# --------------------------------------------------------------------------
# HTML（打印优化，用于在浏览器里"另存为 PDF"）
# --------------------------------------------------------------------------
def _h(text: Any) -> str:
    return html.escape(str(text if text is not None else ""))


def to_html(data: dict) -> str:
    s = data["session"]
    parts: list[str] = []
    parts.append(f"<h1>辩论参谋复盘 — {_h(s['topic'])}</h1>")
    parts.append("<ul class='meta'>")
    parts.append(f"<li><b>我方立场</b>：{_h(s['our_side'])}</li>")
    parts.append(f"<li><b>会话编号</b>：<code>{_h(s['id'])}</code></li>")
    parts.append(f"<li><b>生成时间</b>：{_h(data['generated_at'])}</li>")
    parts.append(
        f"<li><b>对方发言轮次</b>：{data['stats']['turns']}　"
        f"<b>台账条目</b>：{data['stats']['cards']}（成立 {data['stats']['standing']}）</li>"
    )
    parts.append("</ul>")

    parts.append("<h2>一、对方发言记录</h2>")
    if data["turns"]:
        parts.append("<ol class='turns'>")
        for t in data["turns"]:
            parts.append(f"<li>{_h(t['opponent_text'])}</li>")
        parts.append("</ol>")
    else:
        parts.append("<p class='empty'>（无）</p>")

    parts.append("<h2>二、参谋建议</h2>")
    if not data["sections"]:
        parts.append("<p class='empty'>（无）</p>")
    for sec in data["sections"]:
        parts.append(f"<h3>{_h(sec['title'])}</h3>")
        payload = sec.get("payload")
        if sec["advisor"] == "rebutter" and isinstance(payload, list):
            for it in payload:
                parts.append("<div class='card'>")
                parts.append(f"<p class='claim'>{_h(it.get('claim'))}</p>")
                parts.append("<table class='syll'>")
                for label, key in (("大前提", "major_premise"),
                                   ("小前提", "minor_premise"),
                                   ("结　论", "conclusion")):
                    parts.append(
                        f"<tr><th>{label}</th><td>{_h(it.get(key))}</td></tr>"
                    )
                parts.append("</table>")
                parts.append("</div>")
        elif sec["advisor"] == "questioner" and isinstance(payload, list):
            parts.append("<ol class='qs'>")
            for q in payload:
                parts.append(f"<li>{_h(q)}</li>")
            parts.append("</ol>")
        elif sec["advisor"] == "auditor" and isinstance(payload, list):
            if not payload:
                parts.append("<p class='empty'>未发现明显逻辑谬误。</p>")
            parts.append("<ul class='audit'>")
            for it in payload:
                if isinstance(it, dict):
                    parts.append(
                        f"<li><b>{_h(it.get('fallacy'))}</b>｜原话：「{_h(it.get('quote'))}」"
                        f"<br><span class='dim'>{_h(it.get('explain'))}</span></li>"
                    )
            parts.append("</ul>")
        elif sec["advisor"] == "strategist" and isinstance(payload, list):
            if not payload:
                parts.append("<p class='empty'>（未识别到衡量尺度之争）</p>")
            for it in payload:
                if not isinstance(it, dict):
                    continue
                parts.append("<div class='card'>")
                parts.append(
                    f"<p class='claim'>对方用 <b>{_h(it.get('opponent_method'))}</b></p>"
                )
                parts.append(f"<p class='dim'>{_h(it.get('opponent_effect'))}</p>")
                parts.append(
                    f"<p>我方主张 <b>{_h(it.get('our_method'))}</b> 优先："
                    f"{_h(it.get('counter'))}</p>"
                )
                parts.append("</div>")
        elif sec["advisor"] == "risk" and isinstance(payload, list):
            if not payload:
                parts.append("<p class='empty'>（未识别到明显风险）</p>")
            parts.append("<ul class='audit'>")
            for it in payload:
                if isinstance(it, dict):
                    parts.append(
                        f"<li><b>[{_h(it.get('kind'))}]</b> {_h(it.get('risk'))}"
                        f"<br><span class='dim'>应对：{_h(it.get('suggestion'))}</span></li>"
                    )
            parts.append("</ul>")

    parts.append("<h2>三、我方论点台账</h2>")
    if data["cards"]:
        parts.append("<table class='ledger'><thead><tr>"
                     "<th>编号</th><th>主张</th><th>大前提/依据</th><th>状态</th>"
                     "</tr></thead><tbody>")
        for c in data["cards"]:
            status = {"standing": "成立", "weakened": "受损", "abandoned": "放弃"}.get(
                c["status"], c["status"]
            )
            parts.append(
                f"<tr><td><code>{_h(c['id'])}</code></td><td>{_h(c['claim'])}</td>"
                f"<td>{_h(c.get('major_premise') or '—')}</td><td>{_h(status)}</td></tr>"
            )
        parts.append("</tbody></table>")
    else:
        parts.append("<p class='empty'>（无）</p>")

    parts.append("<h2>四、使用提示</h2><ul class='notes'>"
                 "<li>以上内容由 AI 参谋生成，<b>仅作参考</b>，最终判断与取舍由你决定。</li>"
                 f"<li>{NOTE_UNVERIFIED_HTML}</li>"
                 "</ul>")
    parts.append(
        f"<p class='foot'>{_h(_attribution(MAVIS_NAME))}"
        f"　｜　{_h(BASE_NOTE)}</p>"
    )

    body = "\n".join(parts)
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>参谋复盘 — {_h(s['topic'])}</title>
<style>
  :root {{ --ink:#24241f; --dim:#6b6b64; --line:#dcdbd5; }}
  * {{ box-sizing: border-box; }}
  body {{ font-family: -apple-system,"PingFang SC","Songti SC",serif;
         color: var(--ink); max-width: 800px; margin: 0 auto; padding: 40px 28px 64px;
         line-height: 1.75; font-size: 14px; background: #fff; }}
  h1 {{ font-size: 20px; margin: 0 0 16px; }}
  h2 {{ font-size: 16px; margin: 28px 0 10px; padding-bottom: 6px;
        border-bottom: 1px solid var(--line); }}
  h3 {{ font-size: 14px; margin: 18px 0 8px; color: #185fa5; }}
  ul.meta {{ list-style: none; padding: 0; margin: 0 0 8px; font-size: 13px; color: var(--dim); }}
  ul.meta li {{ margin-bottom: 3px; }}
  ul.meta b {{ color: var(--ink); }}
  .card {{ border: 1px solid var(--line); border-radius: 8px; padding: 10px 12px;
           margin-bottom: 10px; break-inside: avoid; }}
  .claim {{ font-weight: 500; margin: 0 0 6px; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
  table.syll th {{ text-align: left; width: 68px; color: var(--dim); font-weight: 400;
                   vertical-align: top; padding: 3px 8px 3px 0; }}
  table.syll td {{ padding: 3px 0; vertical-align: top; }}
  table.ledger th, table.ledger td {{ border: 1px solid var(--line); padding: 5px 8px;
                                       text-align: left; vertical-align: top; }}
  table.ledger th {{ background: #f6f5f1; font-weight: 500; }}
  .empty, .dim {{ color: var(--dim); }}
  code {{ font-family: ui-monospace, Menlo, monospace; font-size: 12px;
          background: #f1efe8; padding: 0 4px; border-radius: 3px; }}
  ol.turns li, ol.qs li, ul.audit li {{ margin-bottom: 6px; }}
  ul.notes li {{ margin-bottom: 4px; font-size: 13px; }}
  .foot {{ margin-top: 20px; padding-top: 10px; border-top: 1px solid var(--line);
           font-size: 12px; color: var(--dim); }}
  .toolbar {{ position: sticky; top: 0; background: #fff; padding: 8px 0 12px;
              border-bottom: 1px solid var(--line); margin-bottom: 8px; }}
  .toolbar button {{ font: inherit; padding: 7px 14px; border: none; border-radius: 8px;
                     background: #185fa5; color: #fff; cursor: pointer; }}
  .toolbar span {{ font-size: 12px; color: var(--dim); margin-left: 10px; }}
  @media print {{ .toolbar {{ display: none; }} body {{ padding: 0; max-width: none; }} }}
</style></head>
<body>
<div class="toolbar">
  <button onclick="window.print()">打印 / 另存为 PDF</button>
  <span>在打印对话框里把目标选为「存储为 PDF」即可；若无对话框或排版异常，请改用较新的浏览器（Chrome / Edge / Firefox）打开本页。</span>
</div>
{body}
</body></html>"""


# --------------------------------------------------------------------------
# Word (.docx)
# --------------------------------------------------------------------------
def to_docx(data: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    s = data["session"]
    doc = Document()

    # 中文字体（东亚洲字体需要显式设置 rFonts，否则 Word 里会回退成宋体/乱码观感）
    style = doc.styles["Normal"]
    style.font.name = "PingFang SC"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        "PingFang SC",
    )

    doc.add_heading(f"辩论参谋复盘 — {s['topic']}", level=1)
    for label, value in (
        ("我方立场", s["our_side"]),
        ("会话编号", s["id"]),
        ("生成时间", data["generated_at"]),
        ("对方发言轮次", str(data["stats"]["turns"])),
        ("台账条目", f"{data['stats']['cards']}（成立 {data['stats']['standing']}）"),
    ):
        p = doc.add_paragraph()
        run = p.add_run(f"{label}：")
        run.bold = True
        p.add_run(str(value))

    doc.add_heading("一、对方发言记录", level=2)
    if data["turns"]:
        for i, t in enumerate(data["turns"], 1):
            doc.add_paragraph(f"{i}. {t['opponent_text']}")
    else:
        doc.add_paragraph("（无）")

    doc.add_heading("二、参谋建议", level=2)
    for sec in data["sections"]:
        doc.add_heading(sec["title"], level=3)
        payload = sec.get("payload")
        if sec["advisor"] == "rebutter" and isinstance(payload, list):
            for i, it in enumerate(payload, 1):
                p = doc.add_paragraph()
                r = p.add_run(f"{i}. {it.get('claim', '')}")
                r.bold = True
                for label, key in (("大前提", "major_premise"),
                                   ("小前提", "minor_premise"),
                                   ("结　论", "conclusion")):
                    doc.add_paragraph(f"{label}：{it.get(key, '')}", style="List Bullet")
        elif sec["advisor"] == "questioner" and isinstance(payload, list):
            for i, q in enumerate(payload, 1):
                doc.add_paragraph(f"{i}. {q}")
        elif sec["advisor"] == "auditor" and isinstance(payload, list):
            if not payload:
                doc.add_paragraph("未发现明显逻辑谬误。")
            for it in payload:
                if isinstance(it, dict):
                    p = doc.add_paragraph()
                    r = p.add_run(it.get("fallacy", "谬误"))
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xA3, 0x2D, 0x2D)
                    p.add_run(f"｜原话：「{it.get('quote', '')}」")
                    doc.add_paragraph(it.get("explain", ""), style="List Bullet")
        elif sec["advisor"] == "strategist" and isinstance(payload, list):
            if not payload:
                doc.add_paragraph("（未识别到衡量尺度之争）")
            for it in payload:
                if isinstance(it, dict):
                    p = doc.add_paragraph()
                    p.add_run("对方用 ").bold = False
                    p.add_run(it.get("opponent_method", "?")).bold = True
                    doc.add_paragraph(it.get("opponent_effect", ""), style="List Bullet")
                    p2 = doc.add_paragraph()
                    p2.add_run("我方主张 ")
                    p2.add_run(it.get("our_method", "?")).bold = True
                    p2.add_run(f" 优先：{it.get('counter', '')}")
        elif sec["advisor"] == "risk" and isinstance(payload, list):
            if not payload:
                doc.add_paragraph("（未识别到明显风险）")
            for it in payload:
                if isinstance(it, dict):
                    p = doc.add_paragraph()
                    r = p.add_run(f"[{it.get('kind', '风险')}]")
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x85, 0x4F, 0x0B)
                    p.add_run(f" {it.get('risk', '')}")
                    doc.add_paragraph(
                        f"应对：{it.get('suggestion', '')}", style="List Bullet"
                    )

    doc.add_heading("三、我方论点台账", level=2)
    if data["cards"]:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ("编号", "主张", "大前提/依据", "状态")):
            cell.text = text
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
        for c in data["cards"]:
            status = {"standing": "成立", "weakened": "受损", "abandoned": "放弃"}.get(
                c["status"], c["status"]
            )
            row = table.add_row().cells
            row[0].text = c["id"]
            row[1].text = c["claim"]
            row[2].text = c.get("major_premise") or "—"
            row[3].text = status
    else:
        doc.add_paragraph("（无）")

    doc.add_heading("四、使用提示", level=2)
    doc.add_paragraph("以上内容由 AI 参谋生成，仅作参考，最终判断与取舍由你决定。")
    doc.add_paragraph(NOTE_UNVERIFIED)

    foot = doc.add_paragraph()
    foot.add_run(_attribution(MAVIS_NAME))
    foot.add_run(f"　｜　{BASE_NOTE}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
